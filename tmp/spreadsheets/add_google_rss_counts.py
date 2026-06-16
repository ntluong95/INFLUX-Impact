from __future__ import annotations

import copy
import json
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

import openpyxl


WORKBOOK_PATH = Path(
    "/Users/luongnguyen/Library/CloudStorage/OneDrive-SharedLibraries-Kungl.Vetenskapsakademien/INFLUX - Documents/3_EPPImpacts/4_LuongImpactWork/INFLUX Impact/data/inputs/EPPs master list.xlsx"
)
WORKSPACE_ROOT = WORKBOOK_PATH.parents[2]
SHEET_NAME = "Emergence review"
NAME_COL = 2
DATA_START_ROW = 3

GROUP_HEADER = "D. RSS search signal"
COLUMN_HEADER = "Google RSS item count (2005-01-01 to 2025-12-31)"
CACHE_PATH = WORKSPACE_ROOT / "tmp" / "spreadsheets" / "google_rss_2005_2025_counts.json"
LOG_PATH = WORKSPACE_ROOT / "tmp" / "spreadsheets" / "google_rss_2005_2025_counts_log.txt"

AFTER_DATE = "2004-12-31"
BEFORE_DATE = "2026-01-01"
USER_AGENT = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36"


def load_cache() -> dict[str, int]:
    if CACHE_PATH.exists():
        return {str(k): int(v) for k, v in json.loads(CACHE_PATH.read_text()).items()}
    return {}


def save_cache(cache: dict[str, int]) -> None:
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2, sort_keys=True))


def build_url(query: str) -> str:
    search = f"{query} after:{AFTER_DATE} before:{BEFORE_DATE}"
    encoded = urllib.parse.quote(search)
    return f"https://news.google.com/rss/search?q={encoded}&hl=en-US&gl=US&ceid=US:en"


def fetch_count(query: str, *, max_attempts: int = 6) -> int:
    url = build_url(query)
    last_error: Exception | None = None
    for attempt in range(1, max_attempts + 1):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(request, timeout=30) as response:
                payload = response.read().decode("utf-8", errors="replace")
                if "rss/unsupported" in getattr(response, "url", ""):
                    return 0
                return len(re.findall(r"<item>", payload))
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if attempt == max_attempts:
                break
            time.sleep(min(20, attempt * 3))
    raise RuntimeError(f"Failed to fetch Google RSS count for {query!r}: {last_error}") from last_error


def ensure_output_column(ws: Any) -> int:
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    for idx, value in enumerate(headers, start=1):
        if value == COLUMN_HEADER:
            return idx

    prior_max_column = ws.max_column
    col_idx = ws.max_column + 1
    ws.cell(row=1, column=col_idx).value = GROUP_HEADER
    ws.cell(row=2, column=col_idx).value = COLUMN_HEADER

    if prior_max_column >= 2:
        for row_idx, source_col in ((1, 14), (2, prior_max_column)):
            source = ws.cell(row=row_idx, column=source_col)
            target = ws.cell(row=row_idx, column=col_idx)
            if source.has_style:
                target._style = copy.copy(source._style)

    width_source = ws.column_dimensions[openpyxl.utils.get_column_letter(prior_max_column)]
    ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(18, width_source.width or 18)

    if ws.auto_filter and ws.auto_filter.ref:
        min_col, min_row, _, max_row = openpyxl.utils.range_boundaries(ws.auto_filter.ref)
        ws.auto_filter.ref = openpyxl.utils.get_column_letter(min_col) + str(min_row) + ":" + openpyxl.utils.get_column_letter(col_idx) + str(ws.max_row)
    return col_idx


def collect_names(ws: Any) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for row in range(DATA_START_ROW, ws.max_row + 1):
        value = ws.cell(row=row, column=NAME_COL).value
        if value is None:
            continue
        name = str(value).strip()
        if not name or name in seen:
            continue
        seen.add(name)
        names.append(name)
    return names


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb[SHEET_NAME]

    output_col = ensure_output_column(ws)
    names = collect_names(ws)
    cache = load_cache()

    started_at = time.time()
    log_lines = [
        f"Workbook: {WORKBOOK_PATH}",
        f"Sheet: {SHEET_NAME}",
        f"Date window: 2005-01-01 to 2025-12-31 (queried as after:{AFTER_DATE} before:{BEFORE_DATE})",
        f"Unique names: {len(names)}",
        "",
    ]

    pending = [name for name in names if name not in cache]
    round_number = 1
    failed: list[str] = []
    while pending:
        print(f"[round] {round_number} pending={len(pending)} cached={len(cache)}")
        failed = []
        for name in pending:
            try:
                cache[name] = fetch_count(name)
                save_cache(cache)
                time.sleep(0.2)
            except Exception as exc:  # noqa: BLE001
                failed.append(name)
                print(f"[retry] {name}: {exc}")
            completed = len(cache)
            if completed % 25 == 0 or completed == len(names):
                elapsed = time.time() - started_at
                print(f"[progress] {completed}/{len(names)} unique names processed in {elapsed:.1f}s")
        if not failed:
            break
        round_number += 1
        if round_number > 5:
            break
        pending = failed
        time.sleep(15)

    if failed:
        raise RuntimeError(f"Unresolved Google RSS queries after retries: {failed}")

    for row in range(DATA_START_ROW, ws.max_row + 1):
        raw_name = ws.cell(row=row, column=NAME_COL).value
        name = str(raw_name).strip() if raw_name is not None else ""
        ws.cell(row=row, column=output_col).value = cache.get(name)

    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    log_lines.extend(f"{name}\t{cache[name]}" for name in names)
    LOG_PATH.write_text("\n".join(log_lines))
    wb.save(WORKBOOK_PATH)

    populated = sum(1 for row in range(DATA_START_ROW, ws.max_row + 1) if ws.cell(row=row, column=output_col).value is not None)
    print(json.dumps({"output_col": output_col, "rows_populated": populated, "unique_names": len(names)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
