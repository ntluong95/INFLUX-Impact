from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(
    "/Users/luongnguyen/Library/CloudStorage/OneDrive-SharedLibraries-Kungl.Vetenskapsakademien/INFLUX - Documents/3_EPPImpacts/4_LuongImpactWork/INFLUX Impact"
)
WORKBOOK = ROOT / "data/inputs/EPPs master list.xlsx"
LOG_PATH = ROOT / "tmp/spreadsheets/plant_emergence_fill_log.txt"


def norm(value: object) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    value = value.replace("\t", " ").strip().lower()
    value = re.sub(r"\s+", " ", value)
    return value


def profile(
    hosts: str,
    first: str,
    reemerge: str,
    endemic: str,
    focal: str,
    suitable: str,
    best_geo: str,
    why: str,
    deprioritize: str,
) -> dict[str, str]:
    return {
        "hosts": hosts,
        "first": first,
        "reemerge": reemerge,
        "endemic": endemic,
        "focal": focal,
        "suitable": suitable,
        "best_geo": best_geo,
        "why": why,
        "deprioritize": deprioritize,
    }


def invasive_plant_profile(hosts: str = "Croplands, rangelands, wetlands, and native plant communities") -> dict[str, str]:
    return profile(
        hosts,
        "Usually native to another region; analytically relevant emergence is first introduction and establishment outside the native range.",
        "Major re-emergence is spread into a new country, watershed, rangeland, or crop zone where the taxon was not previously established.",
        "Native range plus long-established invaded regions where the plant is already naturalized.",
        "Newest high-impact invasion frontier",
        "Maybe",
        "Newest high-impact invasion frontier",
        "Cascade relevance is strongest where invasion alters yields, grazing, water management, biodiversity, or control costs rather than where the species is already background flora.",
        "Long-established invaded areas where the species is now essentially endemic and analytically diffuse.",
    )


def arthropod_pest_profile(hosts: str = "Crops, orchards, plantations, forestry, and horticultural plants") -> dict[str, str]:
    return profile(
        hosts,
        "First described historically in its native range; the key emergence for analysis is first establishment in a new crop-growing or forest region.",
        "Major re-emergence usually reflects trade-mediated spread into a new country or production basin with substantial crop or tree losses.",
        "Native range and invaded regions where the pest is now established.",
        "Recently invaded export-oriented or high-value production region",
        "Yes",
        "Recently invaded export-oriented or high-value production region",
        "Arthropod pests are highly suitable when new establishment triggers quarantine, pesticide shifts, market loss, or visible yield damage.",
        "Areas where the pest has been established for decades and now behaves as a routine endemic problem.",
    )


def forest_pest_profile() -> dict[str, str]:
    return profile(
        "Forest trees, urban trees, and woody hosts",
        "First described in the native forest range; analytically relevant emergence is first destructive establishment outside that range.",
        "Major re-emergence is spread into new forest or urban-tree systems with mortality, sanitation felling, and wood-movement controls.",
        "Native range plus invaded forest regions where the pest/pathogen is established.",
        "Recently invaded forest or urban-tree system",
        "Yes",
        "Recently invaded forest or urban-tree system",
        "Forest EPPs work well for cascade analysis because they alter timber, ecosystem services, fire risk, tourism, and municipal budgets over long time scales.",
        "Long-colonized regions where mortality has stabilized and the signal is no longer clearly emergent.",
    )


def fruit_fly_profile() -> dict[str, str]:
    return profile(
        "Fruit crops and vegetable crops",
        "First described in the native tropical or subtropical range; the key emergence is first incursion into a new fruit-growing region.",
        "Major re-emergence is repeated establishment or outbreak in previously free export-oriented horticultural zones.",
        "Native range and long-established tropical horticultural regions.",
        "Newly invaded fruit-export region",
        "Yes",
        "Newly invaded fruit-export region",
        "Fruit flies are analytically rich because small distribution changes can trigger market closures, eradication campaigns, and major orchard-management costs.",
        "Long-established endemic fruit-fly zones with routine management and limited new signal.",
    )


def plant_pathogen_profile(hosts: str = "Crops, orchards, plantations, nurseries, and forestry hosts") -> dict[str, str]:
    return profile(
        hosts,
        "Usually first described in a native or early affected host range; analytically relevant emergence is first severe disease in a new production system or host-region context.",
        "Major re-emergence commonly reflects nursery trade, planting material movement, vector spread, or first detection in a new country.",
        "Native range and long-established affected production areas.",
        "Newly invaded or high-value production region",
        "Yes",
        "Newly invaded or high-value production region",
        "Plant pathogens are strong cascade candidates where outbreaks reshape production, surveillance, trade, and input use.",
        "Long-endemic settings with chronic background disease and weak event boundaries.",
    )


def plant_bacterial_profile(hosts: str = "Crops, orchards, nurseries, and forestry hosts") -> dict[str, str]:
    return plant_pathogen_profile(hosts)


def plant_fungal_profile(hosts: str = "Crops, orchards, nurseries, and forestry hosts") -> dict[str, str]:
    return plant_pathogen_profile(hosts)


def plant_virus_profile(hosts: str = "Crops and horticultural hosts") -> dict[str, str]:
    return profile(
        hosts,
        "Usually first characterized in a local host range; analytically relevant emergence is first spread into a new crop region or vector-host system.",
        "Major re-emergence often follows trade in propagative material or vector establishment in a previously unaffected production area.",
        "Native/vector-established regions plus places where the virus is now persistent.",
        "Newly invaded crop region with competent vectors or planting-material pathways",
        "Yes",
        "Newly invaded crop region with competent vectors or planting-material pathways",
        "Plant viruses are most useful where new vector-host combinations create rapid management and market shifts.",
        "Long-endemic virus landscapes with weak event separation.",
    )


def plant_nematode_profile(hosts: str = "Crop roots, tubers, bulbs, and perennial hosts") -> dict[str, str]:
    return profile(
        hosts,
        "Usually first described in its original range; the key emergence is first damaging establishment in a new crop-production area.",
        "Major re-emergence is spread through soil, seed, nursery stock, or machinery into previously free production systems.",
        "Native range plus regions with long-established infestation.",
        "Seed-, tuber-, or nursery-intensive production zones",
        "Yes",
        "Seed-, tuber-, or nursery-intensive production zones",
        "Nematodes are analytically valuable where quarantine status, yield loss, and land-use restrictions are strong.",
        "Long-infested regions where the pest is already managed as background risk.",
    )


def invasive_mollusc_profile() -> dict[str, str]:
    return profile(
        "Wetlands, rice systems, aquatic crops, and natural water bodies",
        "First described in its native aquatic range; analytically relevant emergence is establishment in new irrigation or wetland systems.",
        "Major re-emergence is spread into new river basins, wetlands, or crop systems via trade and human-mediated movement.",
        "Native range and long-established invaded aquatic systems.",
        "Newly invaded wetland or irrigated agricultural basin",
        "Maybe",
        "Newly invaded wetland or irrigated agricultural basin",
        "Mollusc invasions matter most where they alter irrigation, aquatic weeds, wetland ecology, or vector/pest management.",
        "Water systems where the species has been established for long periods with weak emergence signal.",
    )


SPECIFIC_PROFILES: dict[str, dict[str, str]] = {
    "xylella fastidiosa": profile(
        "Olive, grapevine, citrus, almond, coffee, and many other woody plants",
        "California, 1892, first recognized in Pierce's disease of grapevine; the bacterium was later defined as Xylella fastidiosa.",
        "Apulia, Italy, 2013 onward, major re-emergence in olive landscapes; additional spread in Europe followed.",
        "The Americas are the long-term endemic range; some European foci are now established.",
        "Apulia, Italy",
        "Yes",
        "Apulia, Italy",
        "Apulia best captures how a newly established xylem-limited bacterium can reshape perennial agriculture, landscapes, and EU phytosanitary policy.",
        "Long-endemic American systems where disease is important but less clearly emergent.",
    ),
    "asian citrus psyllid": profile(
        "Citrus and rutaceous hosts",
        "South and Southeast Asia, first recognized in the native citrus zone.",
        "Major re-emergence occurred through spread to the Americas, especially Brazil and Florida in the 2000s, as the main vector of huanglongbing.",
        "Much of Asia and many citrus regions of the Americas where now established.",
        "Brazil and Florida citrus belts",
        "Yes",
        "Brazil",
        "The psyllid is most analytically relevant where its establishment transforms citrus production through huanglongbing risk and tree decline.",
        "Regions where the insect is long established without a distinct new spread signal.",
    ),
    "trioza erytreae": profile(
        "Citrus and rutaceous hosts",
        "Africa, first recognized in its native range as the African citrus psyllid.",
        "Atlantic islands and the Iberian Peninsula during the 2010s marked the most important recent re-emergence for Europe.",
        "Sub-Saharan Africa and some Atlantic/European invaded areas.",
        "Iberian Peninsula",
        "Yes",
        "Portugal and northwestern Spain",
        "This is the clearest European frontier for citrus-greening vector risk and regional biosecurity response.",
        "Long-endemic African citrus areas when the question is recent spread into new regions.",
    ),
    "xanthomonas citri pv. citri": profile(
        "Citrus",
        "South/Southeast Asia, historically first recognized as citrus canker in the early 20th century.",
        "Repeated re-emergence in the Americas, including Florida and Brazil, with major nursery and orchard control campaigns.",
        "Much of Asia and parts of the Americas where established.",
        "Brazil and Florida",
        "Yes",
        "Brazil",
        "Brazil is the strongest focal geography because canker affects a major export-oriented citrus system with sustained phytosanitary implications.",
        "Long-endemic low-surveillance areas where the disease is routine rather than newly emergent.",
    ),
    "xanthomonas citri pv. aurantifolii": profile(
        "Citrus",
        "South America, first recognized as a distinct canker form affecting citrus.",
        "Regional re-emergence has centered on movement into new citrus-producing areas in Latin America.",
        "Parts of South America.",
        "South American citrus systems",
        "Maybe",
        "South America",
        "Useful mainly where distinct citrus-canker lineages trigger regional quarantine and orchard-management responses.",
        "Areas without active spread or export sensitivity.",
    ),
    "erwinia amylovora": profile(
        "Apple, pear, quince, and other Rosaceae",
        "Northeastern United States, late 18th century, first recognized as fire blight; bacteriology clarified in the late 19th century.",
        "Repeated re-emergence in Europe, the Mediterranean, and other fruit-growing regions through the 20th-21st centuries.",
        "North America and many parts of Europe, the Mediterranean, and western Asia.",
        "Mediterranean and European pome-fruit regions",
        "Yes",
        "Mediterranean Europe",
        "The best current cascade setting is high-value pome-fruit production where nursery restrictions, pruning, and orchard removal scale up quickly.",
        "Long-endemic North American regions where the disease is historically important but less newly emergent.",
    ),
    "pseudomonas syringae pv. actinidiae": profile(
        "Kiwifruit",
        "Japan, 1980s, first recognized kiwifruit bacterial canker due to Psa.",
        "Italy, 2008 onward, and subsequent global spread to major kiwifruit regions including New Zealand.",
        "East Asia and several invaded kiwifruit-producing countries.",
        "New Zealand and Italy",
        "Yes",
        "New Zealand",
        "New Zealand is analytically strongest because a newly spread pathogen affected a nationally important export crop and nursery system.",
        "Regions with minor or historically older low-intensity disease presence.",
    ),
    "ralstonia solanacearum": profile(
        "Potato, tomato, banana, ginger, and many other hosts",
        "Long recognized in tropical regions as bacterial wilt, with the modern species complex clarified over time.",
        "Major re-emergence is repeated introduction into temperate potato systems and new tropical production zones.",
        "Tropical and subtropical regions worldwide; some lineages persist in temperate areas.",
        "Potato seed systems in temperate regions",
        "Yes",
        "Europe and high-value seed-potato systems",
        "The strongest cascade signal appears where a tropical wilt pathogen breaches seed-potato or nursery systems that are expected to remain clean.",
        "Diffuse tropical endemic areas with weak event delimitation.",
    ),
    "ralstonia pseudosolanacearum": plant_bacterial_profile("Tomato, pepper, potato, ornamentals, and other hosts"),
    "ralstonia solanacearum sensu lato": plant_bacterial_profile("Potato, tomato, banana, ginger, and many other hosts"),
    "ralstonia syzygii": profile(
        "Clove and banana-associated systems",
        "Indonesia, long recognized in regional host systems.",
        "Re-emergence is most relevant when the pathogen complex affects new banana or perennial-crop systems in Southeast Asia.",
        "Southeast Asia.",
        "Indonesia and nearby island production systems",
        "Maybe",
        "Indonesia",
        "Its value is highest where perennial-crop decline affects smallholder tree-crop or banana landscapes.",
        "Regions outside the Southeast Asian host-pathogen context.",
    ),
    "clavibacter sepedonicus": profile(
        "Potato",
        "Europe/North America, early 20th century, first recognized as potato ring rot.",
        "Re-emergence occurs via latently infected seed potatoes entering previously free production chains.",
        "Some potato regions in North America and Eurasia; many areas aim to remain free.",
        "Seed-potato production systems",
        "Yes",
        "Northern Europe and North America seed-potato zones",
        "This disease is analytically strongest where a single detection can disrupt seed movement and certification systems.",
        "Areas with chronic low-grade presence and weak trade significance.",
    ),
    "clavibacter michiganensis subsp. michiganensis": plant_bacterial_profile("Tomato and related solanaceous hosts"),
    "dickeya dianthicola": profile(
        "Potato and ornamentals",
        "Europe, historically recognized within soft-rot/blackleg complexes and later distinguished as D. dianthicola.",
        "Major re-emergence has affected European potato seed systems since the 2000s and North America more recently.",
        "Europe and some North American potato systems.",
        "European potato seed systems",
        "Yes",
        "Northwestern Europe",
        "Its strongest cascade effects appear in seed-potato systems where latent infection undermines certification and storage.",
        "Areas without organized seed-potato trade sensitivity.",
    ),
    "acidovorax citrulli": profile(
        "Watermelon, melon, and other cucurbits",
        "United States, 20th century, first recognized as bacterial fruit blotch of watermelon and melon.",
        "Re-emergence has followed seed trade and transplant movement in global cucurbit production.",
        "Many cucurbit-producing regions worldwide.",
        "Seed-intensive cucurbit production systems",
        "Yes",
        "United States and East Asia cucurbit sectors",
        "This disease is most useful where seed transmission drives rapid nursery and fruit-chain impacts.",
        "Regions with little commercial cucurbit seed trade.",
    ),
    "synchytrium endobioticum": profile(
        "Potato",
        "Europe, late 19th century, first recognized as potato wart disease.",
        "Re-emergence matters when new pathotypes are detected in previously free potato districts.",
        "Some potato regions in Europe, Asia, Africa, and the Americas; many areas remain free through regulation.",
        "Potato seed and ware districts with freedom status",
        "Yes",
        "Northern Europe",
        "The disease is analytically valuable where quarantine detections trigger land-use restrictions and long control horizons.",
        "Old infested districts where the signal is no longer clearly emergent.",
    ),
    "fusarium oxysporum f. sp. cubense": profile(
        "Banana and plantain",
        "Panama disease was recognized in the late 19th/early 20th century; the current defining emergence is Tropical Race 4 in Asia.",
        "TR4 re-emerged and spread from Southeast Asia to the Middle East, Africa, and Latin America, especially after 2013-2019.",
        "Long-established in Asia and now in parts of the Middle East, Africa, and Latin America.",
        "TR4 spread into major export banana zones",
        "Yes",
        "Latin American banana frontier and Southeast Asia",
        "TR4 is highly suitable because it threatens clonal monocultures, export earnings, and long-lived plantation systems.",
        "Historic Race 1 landscapes where the disease is an older background problem.",
    ),
    "fusarium circinatum": profile(
        "Pines and other conifers",
        "Southeastern United States, 1940s, first recognized as pitch canker.",
        "Major re-emergence occurred with spread to California, Spain, Portugal, South Africa, and other plantation regions.",
        "Parts of North America and several invaded pine-growing regions worldwide.",
        "Iberian Peninsula and plantation forestry regions",
        "Yes",
        "Iberian Peninsula",
        "This focus best captures nursery pathways, plantation forestry, and longer-term forest-health impacts.",
        "Areas where the pathogen is long established with weak new-spread signal.",
    ),
    "potato late blight agent": profile(
        "Potato and tomato",
        "The Andes/Mexico region is the ancestral center; globally famous emergence occurred in Europe in the 1840s Irish late blight crisis.",
        "Major re-emergence continues through new clonal lineages and fungicide-resistant populations worldwide.",
        "Worldwide in potato and tomato systems.",
        "Ireland 1840s historically; contemporary seed-potato systems for re-emergence",
        "Yes",
        "Historically Ireland; operationally global potato systems",
        "Late blight is one of the clearest crop cascade cases because the historic emergence reshaped food security, migration, and agricultural policy.",
        "Places with low-intensity routine management and no lineage change signal.",
    ),
    "phytophthora kernoviae": profile(
        "Ornamentals, forest trees, and shrubs",
        "United Kingdom/New Zealand, first recognized in the early 2000s as a new Phytophthora disease issue.",
        "Re-emergence is tied to nursery pathways and spread into wider woodland hosts.",
        "Limited foci in the United Kingdom, New Zealand, and a few other regions.",
        "United Kingdom",
        "Yes",
        "United Kingdom",
        "It is most useful where nursery-origin introductions spill into woodland and amenity landscapes.",
        "Regions without confirmed establishment.",
    ),
    "phytophthora lateralis": profile(
        "Chamaecyparis and related conifers",
        "Pacific Northwest, early 20th century, first recognized cedar root/collar disease.",
        "Re-emergence includes spread in North American forests and new incursions into Europe and the UK.",
        "North America and some invaded European areas.",
        "Pacific Northwest and UK nursery/landscape systems",
        "Yes",
        "Pacific Northwest",
        "It best supports cascade analysis where nursery movement and forest decline interact.",
        "Areas with minor or contained detections only.",
    ),
    "sudden oak death agent": profile(
        "Oak, tanoak, ornamentals, and other woody hosts",
        "California, mid-1990s, first recognized as a novel destructive forest and nursery disease caused by Phytophthora ramorum.",
        "Re-emergence has involved repeated nursery-pathway movement and establishment in parts of the UK and Europe.",
        "California/Oregon and several managed/nursery landscapes in Europe.",
        "California",
        "Yes",
        "California",
        "California remains the best focal geography because forest mortality, nursery trade, and urban-wildland interfaces all matter.",
        "Areas with only nursery interceptions or contained outbreaks.",
    ),
    "coffee rust fungus": profile(
        "Coffee",
        "East Africa is the original range; global notoriety followed spread to Asia in the 19th century.",
        "Central America and Colombia, 2012-2014, marked a major modern re-emergence with severe livelihood impacts.",
        "Coffee-growing regions worldwide in the tropics.",
        "Central America 2012-2014",
        "Yes",
        "Central America",
        "This is a strong cascade case because rust affected yields, labor, migration, credit, and renovation costs across smallholder landscapes.",
        "Long-endemic areas with stable management and weaker event definition.",
    ),
    "chestnut blight fungus": profile(
        "Chestnut",
        "Eastern United States, early 1900s, first recognized major introduced blight caused by an Asian-origin fungus.",
        "Re-emergence persists in restoration, orchard, and forest contexts where new spread or hybridization matters.",
        "Asia in native coexistence; North America and Europe in invaded systems.",
        "Eastern United States",
        "Yes",
        "Eastern United States",
        "This remains a classic cascade case because one introduction transformed forest composition, rural economies, and restoration agendas.",
        "Asian native-range systems where the disease is longstanding and not emergent.",
    ),
    "karnal bunt": profile(
        "Wheat and triticale",
        "India, 1931, first recognized as a wheat disease of quarantine significance.",
        "Major re-emergence occurred when detected in Mexico and the United States, where export implications became central.",
        "South Asia and a few additional regions with persistent occurrence.",
        "Export-oriented wheat regions with quarantine sensitivity",
        "Yes",
        "United States and Mexico",
        "Karnal bunt is analytically strongest where even low disease prevalence can drive certification, trade, and surveillance costs.",
        "Long-endemic South Asian areas if the question is specifically new-region emergence.",
    ),
    "fall armyworm": profile(
        "Maize and many other crops",
        "The Americas are the native range; first major Old World emergence was confirmed in West Africa in 2016.",
        "Rapid re-emergence across sub-Saharan Africa, Asia, and Oceania since 2016-2018.",
        "Now established across much of the Americas, Africa, Asia, and Oceania.",
        "Sub-Saharan Africa after 2016",
        "Yes",
        "Sub-Saharan Africa",
        "Africa provides the clearest cascade case because a new invasive pest hit maize-dependent smallholder systems and food-security planning.",
        "The native American range where the pest is long established.",
    ),
    "spodoptera litura": profile(
        "Vegetables, field crops, ornamentals, and other hosts",
        "South and Southeast Asia, first described in the native range.",
        "Re-emergence is most relevant when detected in newly invaded greenhouse or field-crop systems outside its traditional range.",
        "Asia and parts of Oceania/Africa; increasingly invasive elsewhere.",
        "Newly invaded horticultural regions",
        "Maybe",
        "Mediterranean and greenhouse horticultural regions",
        "This pest matters most where rapid spread changes insecticide use and protected-crop biosecurity.",
        "Long-endemic Asian regions with routine management.",
    ),
    "tuta absoluta": profile(
        "Tomato and other solanaceous crops",
        "South America, first described in the native tomato-growing range.",
        "Spain, 2006-2007, marked the defining Old World emergence, followed by spread through the Mediterranean, Africa, and Asia.",
        "South America and now much of the Mediterranean, Africa, and Asia.",
        "Mediterranean Basin",
        "Yes",
        "Mediterranean Basin",
        "This is an excellent cascade case because a new invasive tomato pest rapidly changed greenhouse and open-field production economics.",
        "Native South American range where the pest is older and less sharply emergent.",
    ),
    "tomato leaf miner": profile(
        "Tomato and other solanaceous crops",
        "South America, first described in the native tomato-growing range.",
        "Spain, 2006-2007, marked the defining Old World emergence, followed by spread through the Mediterranean, Africa, and Asia.",
        "South America and now much of the Mediterranean, Africa, and Asia.",
        "Mediterranean Basin",
        "Yes",
        "Mediterranean Basin",
        "This common-name row points to the same invasion dynamic as Tuta absoluta and is most useful in newly invaded tomato systems.",
        "Native South American range where the pest is older and less sharply emergent.",
    ),
    "lycorma delicatula": profile(
        "Grapevine, fruit trees, ornamentals, and many woody hosts",
        "China and surrounding East Asia, first described in the native range.",
        "Major re-emergence includes the Republic of Korea in 2004 and the eastern United States from 2014 onward.",
        "East Asia and the invaded eastern United States.",
        "Eastern United States",
        "Yes",
        "Eastern United States",
        "The U.S. invasion is the clearest cascade case because it affects vineyards, hardwoods, transport corridors, and public response.",
        "Native East Asian range where the insect is less analytically emergent.",
    ),
    "emerald ash borer": profile(
        "Ash trees",
        "Northeast Asia, first described in the native range; destructive emergence outside Asia was first detected in North America in 2002.",
        "North American spread since 2002 and later European concerns are the defining re-emergence story.",
        "Northeast Asia and invaded parts of North America; some European range expansion concerns.",
        "North America after 2002",
        "Yes",
        "North America",
        "This is a strong cascade case because a new wood-borer transformed urban forestry, timber movement rules, and ecosystem services.",
        "Native Asian range where host systems are more co-adapted.",
    ),
    "asian longhorned beetle": profile(
        "Hardwood trees",
        "China and surrounding East Asia, first described in the native range.",
        "Repeated re-emergence via wood-packaging introductions into North America and Europe since the 1990s.",
        "East Asia plus repeatedly detected and sometimes established invaded zones elsewhere.",
        "North America and Europe urban-tree systems",
        "Yes",
        "North America",
        "Urban and peri-urban hardwood landscapes show the clearest cascade effects through sanitation felling, trade, and municipal cost.",
        "Native Asian range where the pest is longstanding.",
    ),
    "anoplophora chinensis": profile(
        "Many woody ornamentals and fruit trees",
        "East Asia, first described in the native range.",
        "Major re-emergence has been repeated introduction into Europe through plant trade since the 2000s.",
        "East Asia and local invaded foci in Europe and elsewhere.",
        "Europe",
        "Yes",
        "Italy and broader Europe",
        "It is most useful where nursery trade creates repeated costly eradication campaigns.",
        "Native Asian range where the insect is less sharply emergent.",
    ),
    "red palm weevil": profile(
        "Date palm, coconut, and ornamental palms",
        "South and Southeast Asia, first recognized in the native range.",
        "Middle East and Mediterranean spread from the 1980s onward represents the major re-emergence.",
        "Asia, the Middle East, North Africa, and much of the Mediterranean where established.",
        "Mediterranean Basin",
        "Yes",
        "Mediterranean Basin",
        "The Mediterranean is the strongest focal geography because the pest altered ornamental and production palms, tourism landscapes, and municipal budgets.",
        "Long-endemic Asian palm systems.",
    ),
    "mediterranean fruit fly": profile(
        "Fruit and vegetable crops",
        "Sub-Saharan Africa is the ancestral range; global emergence followed spread to the Mediterranean and then other continents.",
        "Modern re-emergence is defined by repeated outbreaks in previously free fruit-export regions.",
        "Africa, Mediterranean regions, Latin America, Oceania, and other established areas.",
        "Previously free fruit-export regions",
        "Yes",
        "Mediterranean and outbreak-prone free zones",
        "Medfly is analytically strongest where eradication campaigns and trade sensitivity are immediate.",
        "Long-endemic Mediterranean or African areas where presence is routine.",
    ),
    "oriental fruit fly": fruit_fly_profile(),
    "queensland fruit fly": fruit_fly_profile(),
    "mexican fruit fly": fruit_fly_profile(),
    "caribbean fruit fly": fruit_fly_profile(),
    "west indian fruit fly": fruit_fly_profile(),
    "natal fruit fly": fruit_fly_profile(),
    "olive fruit fly": fruit_fly_profile(),
    "oriental citrus fly": fruit_fly_profile(),
    "bactrocera latifrons": fruit_fly_profile(),
    "bactrocera kandiensis": fruit_fly_profile(),
    "bactrocera occipitalis": fruit_fly_profile(),
    "bactrocera pyrifoliae": fruit_fly_profile(),
    "anastrepha fraterculus": fruit_fly_profile(),
    "drosophila suzukii": profile(
        "Soft fruit, cherry, berry, and stone-fruit crops",
        "East Asia, first described in the native range.",
        "North America and Europe, 2008 onward, major re-emergence in soft-fruit systems.",
        "East Asia and now much of Europe and the Americas where established.",
        "Europe and North America soft-fruit sectors",
        "Yes",
        "Europe",
        "This pest is analytically strong because newly invasive oviposition in ripening fruit changed harvest, labor, and spray regimes quickly.",
        "Native East Asian range where the pest is long established.",
    ),
    "colorado potato beetle": profile(
        "Potato and other solanaceous crops",
        "North America, 19th century, first recognized as a severe potato pest.",
        "Repeated re-emergence in Europe and Asia reflects insecticide resistance and new spread into potato systems.",
        "North America and much of Europe/Asia where established.",
        "Europe and Eurasian potato systems",
        "Yes",
        "Europe",
        "The best cascade setting is where resistance and spread raise control costs in intensive potato production.",
        "North America where the beetle is long established background pest.",
    ),
    "western corn rootworm": profile(
        "Maize",
        "North America, first recognized in the native maize-growing range.",
        "Europe since the 1990s represented the defining new-region emergence.",
        "North America and some established European areas.",
        "Central and southeastern Europe",
        "Yes",
        "Central Europe",
        "This is a clean crop-pest invasion case affecting maize rotations, soil insecticides, and biosecurity policy.",
        "North American native range where it is longstanding.",
    ),
    "northern corn rootworm": arthropod_pest_profile("Maize"),
    "desert locust": profile(
        "Rainfed and irrigated field crops, rangeland vegetation, and pasture",
        "Ancient transboundary pest of arid Africa and Southwest Asia rather than a new-to-science emergence.",
        "Major re-emergence occurred in the Horn of Africa, Arabian Peninsula, and Southwest Asia in 2019-2021.",
        "Recession areas across North Africa, the Sahel, the Horn of Africa, the Arabian Peninsula, and Southwest Asia.",
        "Horn of Africa 2019-2021",
        "Yes",
        "Horn of Africa",
        "That upsurge is analytically rich because it linked climate anomalies, food security, pastoral livelihoods, and emergency control operations.",
        "Recession areas with only low-density background presence.",
    ),
    "giant salvinia": profile(
        "Freshwater bodies, rice systems, canals, and wetlands",
        "South America, native range; emergence elsewhere is via first establishment as an invasive floating fern.",
        "Major re-emergence has occurred through spread in tropical and subtropical wetlands in Africa, Asia, Oceania, and the Americas.",
        "South America and many invaded warm freshwater systems worldwide.",
        "Newly invaded wetland and irrigation systems",
        "Yes",
        "Tropical wetland and irrigation frontiers",
        "This invasion matters where water control, fisheries, rice systems, and biodiversity are disrupted simultaneously.",
        "Long-infested water bodies where the species is now background management burden.",
    ),
    "pig weed, alligator weed, alligatorweed": profile(
        "Aquatic systems, irrigated crops, and riparian habitats",
        "South America, native range; emergence elsewhere is via invasive establishment.",
        "Major re-emergence has occurred through spread in waterways and irrigated agriculture in Asia, Australasia, and North America.",
        "South America plus many invaded warm regions.",
        "Irrigated and aquatic invasion frontiers",
        "Yes",
        "China and Australia-type irrigation landscapes",
        "Alligator weed is useful where a single invasive plant affects water delivery, aquatic control, and farm margins at once.",
        "Long-established invaded areas with routine background management.",
    ),
    "witchweed": profile(
        "Cereal and grass crops, especially maize and sorghum",
        "Africa, native range; emergence elsewhere is via first establishment in new cereal systems.",
        "Major re-emergence occurred with spread into the United States and continuing severe impacts in African smallholder maize and sorghum systems.",
        "Sub-Saharan Africa plus limited introduced regions.",
        "Sub-Saharan African cereal belts",
        "Yes",
        "Sub-Saharan Africa",
        "Witchweed is analytically valuable where parasitic weed pressure directly undermines staple-crop food security and labor demand.",
        "Areas where it is absent or tightly contained.",
    ),
    "purple witchweed": profile(
        "Cereal and grass crops, especially maize and sorghum",
        "Africa, native range; emergence elsewhere is via first establishment in new cereal systems.",
        "Major re-emergence has occurred through spread into new cereal-growing zones and persistent African smallholder systems.",
        "Sub-Saharan Africa and scattered invaded regions.",
        "Sub-Saharan African cereal belts",
        "Yes",
        "Sub-Saharan Africa",
        "Parasitic weeds create strong cascade effects in staple-crop systems with few affordable control options.",
        "Areas where the weed is long established without a distinct new spread signal.",
    ),
    "tobacco witchweed": invasive_plant_profile("Tobacco, maize, sugarcane, and other crop systems"),
    "kudzu vine": profile(
        "Forestry edges, transport corridors, orchards, and field margins",
        "East Asia, native range; emergence elsewhere is invasive establishment.",
        "Major re-emergence in the southeastern United States and newer spread into additional temperate regions.",
        "East Asia and long-invaded parts of the southeastern United States.",
        "North American invasion frontiers",
        "Maybe",
        "Southeastern United States",
        "Kudzu matters most where rapid vegetative spread changes land use, utility management, and invasion-control costs.",
        "Areas where it has been long naturalized with stable management burdens.",
    ),
    "pueraria montana": profile(
        "Forestry edges, transport corridors, orchards, and field margins",
        "East Asia, native range; emergence elsewhere is invasive establishment.",
        "Major re-emergence in the southeastern United States and newer spread into additional temperate regions.",
        "East Asia and long-invaded parts of the southeastern United States.",
        "North American invasion frontiers",
        "Maybe",
        "Southeastern United States",
        "This taxon is most useful where vine invasion changes land-cover management and transport/utility costs.",
        "Areas where it is long naturalized and analytically diffuse.",
    ),
    "water lettuce": profile(
        "Freshwater bodies, canals, wetlands, and rice/aquatic systems",
        "Tropical to subtropical native range debated; analytic emergence elsewhere is invasive establishment in water systems.",
        "Major re-emergence is spread through warm waterways and irrigation systems beyond prior range.",
        "Tropical and subtropical regions where established.",
        "Newly invaded freshwater and irrigation systems",
        "Maybe",
        "Warm irrigation and wetland systems",
        "It is useful where floating-mat invasion disrupts water management, fisheries, and aquatic vegetation control.",
        "Long-infested systems where the plant is already background management burden.",
    ),
    "pomacea maculata": invasive_mollusc_profile(),
    "channeled applesnail": invasive_mollusc_profile(),
    "pine wood nematode": profile(
        "Pine and other conifers",
        "North America, first described in the native range; devastating pine wilt emergence outside the native range was later seen in Asia.",
        "Japan and then China, Korea, Portugal, and Spain represent the major re-emergence pathway.",
        "North America native range and established invaded foci in Asia and Europe.",
        "East Asia and Iberia",
        "Yes",
        "East Asia",
        "Pine wilt is analytically strong because new nematode-vector establishment can alter forestry, timber movement, and landscape mortality.",
        "North American native range where hosts are more co-adapted.",
    ),
    "globodera pallida": profile(
        "Potato",
        "Andean South America, native range; emergence elsewhere is invasion of potato systems through tuber and soil movement.",
        "Major re-emergence has occurred in Europe and other seed-potato systems beyond the Andes.",
        "Andean region plus established potato-growing regions in Europe and elsewhere.",
        "Seed-potato systems outside the Andes",
        "Yes",
        "Northern Europe",
        "This nematode is most useful where quarantine status and resistant cultivar choices shape potato-sector costs.",
        "Andean native range where the pest is longstanding.",
    ),
    "globodera rostochiensis": profile(
        "Potato",
        "Andean South America, native range; emergence elsewhere is invasion of potato systems through tuber and soil movement.",
        "Major re-emergence has occurred in Europe and other seed-potato systems beyond the Andes.",
        "Andean region plus established potato-growing regions in Europe and elsewhere.",
        "Seed-potato systems outside the Andes",
        "Yes",
        "Northern Europe",
        "This species is analytically strongest where quarantine detections affect potato rotations and trade.",
        "Andean native range where the pest is longstanding.",
    ),
    "meloidogyne enterolobii": plant_nematode_profile("Vegetables, fruit trees, and nursery crops"),
    "meloidogyne chitwoodi": plant_nematode_profile("Potato, vegetables, and field crops"),
    "xiphinema californicum": plant_nematode_profile("Vineyards, orchards, and perennial crops"),
    "xiphinema rivesi": plant_nematode_profile("Vineyards, orchards, and perennial crops"),
    "cucumber mosaic cucumovirus": plant_virus_profile("Vegetables, ornamentals, and many broadleaf crops"),
    "pseudomonas syringae pv. tomato": plant_bacterial_profile("Tomato"),
    "pseudomonas syringae pv. actinidifoliorum": plant_bacterial_profile("Kiwifruit"),
    "xanthomonas fragariae": plant_bacterial_profile("Strawberry"),
    "xylophilus ampelinus": plant_bacterial_profile("Grapevine"),
    "phyllosticta citricarpa": plant_fungal_profile("Citrus"),
    "phytophthora fragariae": plant_fungal_profile("Strawberry"),
    "phytophthora rubi": plant_fungal_profile("Raspberry and blackberry"),
    "verticillium dahliae": plant_fungal_profile("Many crops and woody hosts"),
    "verticillium nonalfalfae": plant_fungal_profile("Hop, tree hosts, and other plants"),
    "monilinia fructicola": plant_fungal_profile("Stone fruit"),
    "bretziella fagacearum": forest_pest_profile(),
    "geosmithia morbida": forest_pest_profile(),
    "ceratocystis platani": forest_pest_profile(),
    "heterobasidion irregulare": forest_pest_profile(),
}

PLANT_GENERA = {
    "abutilon", "acanthospermum", "acroptilon", "amaranthus", "ambrosia",
    "calystegia", "cenchrus", "commelina", "cuscuta", "cyperus", "digitaria",
    "echinochloa", "eleusine", "euphorbia", "gymnocoronis", "hakea",
    "heracleum", "hordeum", "humulus", "ipomoea", "lespedeza", "panicum",
    "parthenium", "pueraria", "salvinia", "sesbania", "sicyos", "solanum",
    "xanthium",
}

ARTHROPOD_GENERA = {
    "acrobasis", "aculops", "agrilus", "aleurocanthus", "anoplophora",
    "aphis", "bactericera", "bactrocera", "bemisia", "cacopsylla",
    "cadra", "chilo", "choristoneura", "chrysodeixis", "coccus",
    "conotrachelus", "cydia", "dasineura", "dialeurodes", "diuraphis",
    "eriosoma", "frankliniella", "grapholita", "helicoverpa", "liriomyza",
    "lissorhoptrus", "maconellicoccus", "manduca", "obolodiplosis",
    "ophiomyia", "panonychus", "pectinophora", "pentalonia", "phenacoccus",
    "phthorimaea", "pityophthorus", "pseudococcus", "pterochloroides",
    "pulvinaria", "rhagoletis", "saperda", "scirtothrips", "sipha",
    "zeugodacus", "trogoderma", "crisicoccus", "epitrix", "nemorimyza",
    "phyllocoptes", "oligonychus", "trioza", "lycorma", "thrips",
    "sternochetus", "rhynchophorus", "tetranychus", "platynota",
    "leucinodes", "gymnandrosoma", "prodiplosis", "thaumatotibia", "tuta",
    "daktulosphaira",
}


def infer_hosts(name: str, sci: str) -> str:
    text = f"{name} {sci}"
    rules = [
        (("citrus", "orange", "mandarin", "grapefruit", "lime"), "Citrus"),
        (("potato", "tuber"), "Potato and other solanaceous crops"),
        (("tomato",), "Tomato and other solanaceous crops"),
        (("banana", "plantain"), "Banana and plantain"),
        (("coffee",), "Coffee"),
        (("rice",), "Rice"),
        (("maize", "corn"), "Maize"),
        (("wheat",), "Wheat"),
        (("apple",), "Apple and pome fruit"),
        (("pear",), "Pear and pome fruit"),
        (("kiwi", "actinid"), "Kiwifruit"),
        (("strawberry",), "Strawberry"),
        (("grape", "vine"), "Grapevine and vineyard hosts"),
        (("pine", "spruce", "fir", "hemlock", "ash", "oak", "cedar", "walnut", "chestnut", "elm"), "Forest and urban trees"),
        (("fruit fly", "berry", "fruit"), "Fruit crops"),
        (("weevil", "aphid", "psyllid", "whitefly", "mite", "thrips"), "Crops, orchards, forestry, and horticultural plants"),
        (("ragweed", "spurge", "nightshade", "grass", "weed", "fern", "lettuce", "water"), "Croplands, rangelands, wetlands, and native plant communities"),
    ]
    for tokens, hosts in rules:
        if any(token in text for token in tokens):
            return hosts
    return "Plants and crop or ecosystem hosts"


def sci_specific_key(sci: str) -> str:
    sci_key = norm(sci)
    sci_key = re.sub(r"\s+\[.*\]$", "", sci_key)
    return sci_key


def fallback_profile(name: str, sci: str) -> tuple[dict[str, str], str]:
    text = f"{name} {sci}"
    hosts = infer_hosts(name, sci)
    genus = sci.split(" ", 1)[0] if sci else ""

    if any(token in text for token in [
        "virus", "viroid", "phytoplasma", "cucumovirus"
    ]):
        return plant_virus_profile(hosts), "viral"

    if any(token in text for token in [
        "xanthomonas", "pseudomonas", "ralstonia", "erwinia", "clavibacter",
        "dickeya", "pantoea", "xylella", "xylophilus", "curtobacterium",
        "paraburkholderia", "acidovorax"
    ]):
        return plant_bacterial_profile(hosts), "bacterial"

    if any(token in text for token in [
        "phytophthora", "fusarium", "alternaria", "ceratocystis", "diaporthe",
        "monilinia", "verticillium", "puccinia", "synchytrium", "cronartium",
        "gymnosporangium", "phialophora", "phyllosticta", "phymatotrichopsis",
        "stagonosporopsis", "sphaerulina", "thekopsora", "glomerella",
        "pseudocercospora", "stegophora", "stenocarpella", "black knot",
        "rust", "bunt", "blight", "smut", "mildew", "canker", "rot", "fungus",
        "agent"
    ]):
        if any(token in text for token in ["oak", "pine", "spruce", "fir", "hemlock", "ash", "elm", "walnut", "chestnut"]):
            return forest_pest_profile(), "forest_pathogen"
        return plant_fungal_profile(hosts), "fungal"

    if any(token in text for token in [
        "nematode", "globodera", "meloidogyne", "aphelenchoides",
        "ditylenchus", "bursaphelenchus", "xiphinema", "radopholus",
        "nacobbus"
    ]):
        return plant_nematode_profile(hosts), "nematode"

    if any(token in text for token in ["snail", "pomacea", "applesnail"]):
        return invasive_mollusc_profile(), "mollusc"

    if any(token in text for token in [
        "fruit fly", "bactrocera", "anastrepha", "ceratitis",
        "drosophila suzukii", "walnut husk fly", "olive fruit fly",
        "rhagoletis", "apple maggot", "zeugodacus"
    ]):
        return fruit_fly_profile(), "fruit_fly"

    if any(token in text for token in [
        "beetle", "borer", "sawyer", "engraver", "twig borer",
        "weevil", "rootworm", "bark beetle", "sirex", "agrilus",
        "monochamus", "ips ", "pissodes", "polygraphus", "xylotrechus",
        "tetropium", "trichoferus", "euwallacea"
    ]):
        if any(token in text for token in ["pine", "spruce", "fir", "hemlock", "ash", "oak", "elm", "walnut", "chestnut"]):
            return forest_pest_profile(), "forest_insect"
        return arthropod_pest_profile(hosts), "arthropod"

    if genus in ARTHROPOD_GENERA:
        if any(token in text for token in ["pine", "spruce", "fir", "hemlock", "ash", "oak", "elm", "walnut", "chestnut"]):
            return forest_pest_profile(), "forest_insect"
        return arthropod_pest_profile(hosts), "arthropod_genus"

    if any(token in text for token in [
        "aphid", "psyllid", "whitefly", "mite", "thrips", "leafhopper",
        "mealybug", "scale", "stinkbug", "planococcus", "ripersiella"
    ]):
        return arthropod_pest_profile(hosts), "sap_feeder"

    if any(token in text for token in [
        "moth", "armyworm", "leafminer", "leafroller", "budworm",
        "cutworm", "looper", "caterpillar", "locust", "borer", "fly"
    ]):
        return arthropod_pest_profile(hosts), "arthropod"

    if any(token in text for token in [
        "ragweed", "spurge", "nightshade", "grass", "weed", "fern",
        "knapweed", "bindweed", "dodder", "sedge", "primrose",
        "witchweed", "pampas", "water lettuce", "humulus", "heracleum",
        "lespedeza", "kudzu", "salvinia"
    ]):
        return invasive_plant_profile(hosts), "invasive_plant"

    if genus in PLANT_GENERA:
        return invasive_plant_profile(hosts), "plant_taxon"

    if sci and " " in sci and not any(token in sci for token in [". pv.", "f. sp.", "subsp.", "virus", "viroid"]):
        return invasive_plant_profile(hosts), "plant_taxon"

    return plant_pathogen_profile(hosts), "generic"


def profile_for_row(name: str, sci: str) -> tuple[dict[str, str], str]:
    name_key = norm(name)
    sci_key = sci_specific_key(sci)
    if name_key in SPECIFIC_PROFILES:
        return SPECIFIC_PROFILES[name_key], f"name_profile:{name_key}"
    if sci_key in SPECIFIC_PROFILES:
        return SPECIFIC_PROFILES[sci_key], f"sci_profile:{sci_key}"
    result, mode = fallback_profile(name_key, sci_key)
    return result, f"fallback:{mode}"


def main() -> None:
    wb = load_workbook(WORKBOOK)
    review = wb["Emergence review"]
    fallback_rows: list[str] = []

    for row in range(3, review.max_row + 1):
        if norm(review.cell(row, 5).value) != "plant":
            continue

        name = str(review.cell(row, 2).value or "")
        sci = str(review.cell(row, 3).value or "")
        result, mode = profile_for_row(name, sci)

        if mode.startswith("fallback:"):
            fallback_rows.append(f"{row}\t{name}\t{sci}\t{mode}")

        review.cell(row, 4).value = result["hosts"]
        review.cell(row, 6).value = result["first"]
        review.cell(row, 7).value = result["reemerge"]
        review.cell(row, 8).value = result["endemic"]
        review.cell(row, 9).value = result["focal"]
        review.cell(row, 10).value = result["suitable"]
        review.cell(row, 11).value = result["best_geo"]
        review.cell(row, 12).value = result["why"]
        review.cell(row, 13).value = result["deprioritize"]

    wb.save(WORKBOOK)

    lines = [
        "Plant emergence fill log",
        f"Workbook: {WORKBOOK}",
        f"Fallback rows: {len(fallback_rows)}",
        "",
    ]
    lines.extend(fallback_rows)
    LOG_PATH.write_text("\n".join(lines))
    print(f"Saved workbook: {WORKBOOK}")
    print(f"Wrote log: {LOG_PATH}")
    print(f"Fallback rows: {len(fallback_rows)}")


if __name__ == "__main__":
    main()
