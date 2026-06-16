from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from epp_emergence_analysis_config import EVENT_DEFINITION, STANDARD_NAME_COLUMN, yes_no


def save_json(payload: dict[str, Any], output_dir: Path, stem: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    return write_json(payload, output_dir / f"{stem}.json")


def write_json(payload: dict[str, Any], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_excel_from_payload(payload: dict[str, Any], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    readme_rows = [
        {"Field": "Workbook", "Value": "OpenAI EPP emergence event output"},
        {"Field": "Model", "Value": payload.get("metadata", {}).get("model", "")},
        {"Field": "Event definition", "Value": EVENT_DEFINITION},
        {"Field": "Generated at UTC", "Value": payload.get("metadata", {}).get("generated_at_utc", "")},
    ]
    epp_rows, event_rows, source_rows = flatten_results(payload)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(readme_rows).to_excel(writer, sheet_name="README", index=False)
        pd.DataFrame(event_rows).to_excel(writer, sheet_name="Event sheet", index=False)
        pd.DataFrame(epp_rows).to_excel(writer, sheet_name="EPP sheet", index=False)
        pd.DataFrame(payload.get("selected_rows", [])).to_excel(writer, sheet_name="Selected inputs", index=False)
        pd.DataFrame(payload.get("country_rows", [])).to_excel(writer, sheet_name="Country", index=False)
        pd.DataFrame(source_rows).to_excel(writer, sheet_name="Source_URLs", index=False)
        pd.DataFrame(payload.get("warnings", [])).to_excel(writer, sheet_name="Run warnings", index=False)
    style_workbook(output_path)
    return output_path


def flatten_results(payload: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    epp_rows: list[dict[str, Any]] = []
    event_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, Any]] = []
    for result in payload.get("results", []):
        input_row = result.get("input_row", {})
        parsed = result.get("parsed_response", {})
        epp_name = parsed.get("scientific_name_standardize") or input_row.get(STANDARD_NAME_COLUMN, "")
        epp_rows.append(epp_summary_row(input_row, parsed, result, epp_name))
        for source in parsed.get("source_urls", []):
            source_rows.append({"Input row": input_row.get("Input row", ""), "EPP scientific name": epp_name, **source})
        for event in parsed.get("events", []):
            event_rows.append(event_sheet_row(input_row, parsed, event, epp_name))
    return epp_rows, event_rows, source_rows


def epp_summary_row(input_row: dict[str, Any], parsed: dict[str, Any], result: dict[str, Any], epp_name: str) -> dict[str, Any]:
    return {
        "Input row": input_row.get("Input row", ""),
        "EPP scientific name": epp_name,
        "EPP common name": parsed.get("common_name", input_row.get("Common name", "")),
        "Main host system": parsed.get("host_system", input_row.get("Host system", "")),
        "EPP type": input_row.get("EPP type", ""),
        "Emergence suitability": parsed.get("emergence_suitability", ""),
        "First emergence time": parsed.get("first_emergence_time", ""),
        "First emergence place": parsed.get("first_emergence_place", ""),
        "First emergence summary": parsed.get("first_emergence_summary", ""),
        "Major sub-regions 2005-2025": parsed.get("major_subregions_2005_2025", ""),
        "Emergence event count 2005-2025": parsed.get("emergence_event_count_2005_2025", ""),
        "No event reason": parsed.get("no_event_reason", ""),
        "Quality notes": parsed.get("quality_notes", ""),
        "OpenAI response ID": result.get("openai_response_id", ""),
    }


def event_sheet_row(input_row: dict[str, Any], parsed: dict[str, Any], event: dict[str, Any], epp_name: str) -> dict[str, Any]:
    return {
        "Input row": input_row.get("Input row", ""),
        "EPP scientific name": epp_name,
        "EPP common name": parsed.get("common_name", input_row.get("Common name", "")),
        "Year": event.get("year", ""),
        "Continent": event.get("continent", ""),
        "Sub-regions": event.get("sub_regions", ""),
        "Countries": event.get("countries", ""),
        "Host": event.get("host", ""),
        "First/global emergence": yes_no(event.get("first_global_emergence")),
        "First detection": yes_no(event.get("first_detection")),
        "Intra-continental emergence": yes_no(event.get("intra_continental_emergence")),
        "Inter-continental emergence": yes_no(event.get("inter_continental_emergence")),
        "Re-occurrence": yes_no(event.get("re_occurrence")),
        "Re-emergence": yes_no(event.get("re_emergence")),
        "Confidence": event.get("confidence", ""),
        "Event summary": event.get("event_summary", ""),
        "Mechanism or driver": event.get("mechanism_or_driver", ""),
        "Source URLs": "; ".join(event.get("source_urls", [])),
        "Review": event.get("review", ""),
    }


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx, column_cells in enumerate(sheet.columns, start=1):
            max_len = max(len(str(cell.value or "")) for cell in column_cells[:100])
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 60)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)
